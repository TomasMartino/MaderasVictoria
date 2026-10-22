"""One-off data repair: row 394 holds a date as text ("1/06//2026")."""
import datetime
import openpyxl

P = r"C:\Users\User\Desktop\Proyecos\TioVity\Ingresos Y Gastos Semanales (1).xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["VENTAS DIARIAS"]

for r in range(3, ws.max_row + 1):
    v = ws.cell(r, 2).value
    if isinstance(v, str) and v.strip():
        print("fila", r, "->", repr(v), "| vecinos:", repr(ws.cell(r - 1, 2).value), repr(ws.cell(r + 1, 2).value))
        fmt = ws.cell(r - 1, 2).number_format
        ws.cell(r, 2).value = datetime.datetime(2026, 6, 1)
        ws.cell(r, 2).number_format = fmt
        print("   corregido a 01/06/2026 con formato", repr(fmt))

wb.save(P)

wb2 = openpyxl.load_workbook(P)
ws2 = wb2["VENTAS DIARIAS"]
dates = [ws2.cell(r, 2).value for r in range(3, ws2.max_row + 1)
         if isinstance(ws2.cell(r, 2).value, datetime.datetime)]
print("fechas validas:", len(dates), "| min", min(dates).date(), "| max", max(dates).date())
oct_ = sum(ws2.cell(r, 5).value for r in range(3, ws2.max_row + 1)
           if isinstance(ws2.cell(r, 2).value, datetime.datetime)
           and ws2.cell(r, 2).value.month == 10 and isinstance(ws2.cell(r, 5).value, (int, float)))
print("total octubre:", oct_)
print("balance C7:", wb2["BALANCE"]["C7"].value)
