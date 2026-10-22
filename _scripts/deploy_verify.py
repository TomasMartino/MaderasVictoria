"""Post-deploy check: verify the files now sitting in the user's folder."""
import os
import zipfile

import openpyxl

BASE = r"C:\Users\User\Desktop\Proyecos\TioVity"
PRESU = os.path.join(BASE, "PRESUPUESTO.xlsx")
INGRE = os.path.join(BASE, "Ingresos Y Gastos Semanales (1).xlsx")
LISTA = os.path.join(BASE, "LISTA DE PRECIOS.xlsx")
BAK = os.path.join(BASE, "respaldo 2026-10-22")

ok, fail = [], []


def check(c, m):
    (ok if c else fail).append(m)


wb = openpyxl.load_workbook(PRESU)
check(wb.sheetnames == ["CLIENTES", "PRECIOS"], str(wb.sheetnames))
check(wb["PRECIOS"].sheet_state == "hidden", "PRECIOS oculta")
check(wb["CLIENTES"]["C8"].value is None, "presupuesto limpio (sin restos de prueba)")
check(wb["CLIENTES"]["L86"].value == "=SUM(M8:M85)", "total")
check(len(wb["PRECIOS"]["A:A"]) > 80, "catalogo cargado")
with zipfile.ZipFile(PRESU) as z, zipfile.ZipFile(os.path.join(BAK, "PRESUPUESTO.xlsx")) as zb:
    check(z.read("xl/media/image1.png") == zb.read("xl/media/image1.png"), "logo intacto")

wb2 = openpyxl.load_workbook(INGRE)
check(str(wb2["BALANCE"]["C7"].value).startswith("=IF($A7"), "balance automatizado")
check(wb2["BALANCE"]["C39"].value == "=C10+C17+C24+C31+C38", "total mensual")
check(wb2["BALANCE"]["F38"].value == "=SUM(F32:F37)", "bug del total semanal corregido")
check(wb2["VENTAS DIARIAS"].max_row <= 500, f"ventas sin filas de prueba: {wb2['VENTAS DIARIAS'].max_row}")

wb3 = openpyxl.load_workbook(LISTA)
check(wb3["Hoja1"]["B89"].value == "TODOS LOS PRECIOS SON MAS IVA", "lista completa")

print(f"OK {len(ok)} / FALLAS {len(fail)}")
for f in fail:
    print(" X", f)
