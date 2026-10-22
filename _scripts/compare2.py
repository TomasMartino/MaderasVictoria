"""Accurate diff: compare values, formats and RESOLVED visual styles (not internal ids)."""
import openpyxl

BAK = r"C:\Users\User\Desktop\Proyecos\TioVity\respaldo 2026-10-22\Ingresos Y Gastos Semanales (1).xlsx"
NEW = r"C:\Users\User\Desktop\Proyecos\TioVity\Ingresos Y Gastos Semanales (1).xlsx"

a = openpyxl.load_workbook(BAK)
b = openpyxl.load_workbook(NEW)


def col(c):
    return getattr(c, "rgb", None) or str(getattr(c, "theme", "")) + "/" + str(getattr(c, "tint", ""))


def sig(cell):
    f, fl, bd, al = cell.font, cell.fill, cell.border, cell.alignment
    return (
        cell.value,
        cell.number_format,
        (f.name, f.sz, bool(f.b), bool(f.i), col(f.color) if f.color else None),
        (fl.fill_type, col(fl.fgColor) if fl.fill_type else None),
        tuple((getattr(s, "style", None), getattr(getattr(s, "color", None), "rgb", None))
              for s in (bd.left, bd.right, bd.top, bd.bottom)),
        (al.horizontal, al.vertical, bool(al.wrap_text)),
    )


for name in a.sheetnames:
    wa, wb_ = a[name], b[name]
    diffs = []
    mr = max(wa.max_row, wb_.max_row)
    mc = max(wa.max_column, wb_.max_column)
    for row in wa.iter_rows(min_row=1, max_row=mr, max_col=mc):
        for c in row:
            if sig(c) != sig(wb_[c.coordinate]):
                diffs.append((c.coordinate, c.value, wb_[c.coordinate].value))
    print(f"=== {name} === celdas distintas: {len(diffs)}")
    for d in diffs[:50]:
        print("   ", d)
    if len(diffs) > 50:
        print(f"    ... y {len(diffs)-50} mas")
