"""Build the automated PRESUPUESTO, BALANCE and LISTA DE PRECIOS workbooks.

Writes to ./out (never in place). Usage: python build.py
"""
import os
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import catalog_gen as catmod

BASE = r"C:\Users\User\Desktop\Proyecos\TioVity"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
# always build from the pristine originals so the script is re-runnable
SRC = os.path.join(BASE, "respaldo 2026-10-22")
LISTA = os.path.join(SRC, "LISTA DE PRECIOS.xlsx")
PRESU = os.path.join(SRC, "PRESUPUESTO.xlsx")
INGRE = os.path.join(SRC, "Ingresos Y Gastos Semanales (1).xlsx")

MONEY = '[$$]#,##0'
DETAIL_FIRST, DETAIL_LAST = 8, 85

catalog = catmod.catalog
LAST_CAT = len(catalog) + 1          # last row of the catalog sheet
CATALOG_NAMES = [c["name"] for c in catalog]

os.makedirs(OUT, exist_ok=True)


# ---------------------------------------------------------------- PRESUPUESTO
def build_presupuesto():
    wb = openpyxl.load_workbook(PRESU)
    ws = wb["CLIENTES"]

    # --- 1. repair the detail grid (rows 26-53 lost their price columns)
    for mr in list(ws.merged_cells.ranges):
        if str(mr) == "G26:K53":
            ws.unmerge_cells("G26:K53")
    for r in list(range(26, 54)):
        for a, b in (("A", "B"), ("C", "F"), ("H", "I"), ("J", "K")):
            if f"{a}{r}:{b}{r}" not in [str(m) for m in ws.merged_cells.ranges]:
                ws.merge_cells(f"{a}{r}:{b}{r}")
    # rows 84-85 had precio columns merged across two rows
    for mr in list(ws.merged_cells.ranges):
        if str(mr) in ("M84:M85", "J84:L85"):
            ws.unmerge_cells(str(mr))
    for r in (84, 85):
        for a, b in (("A", "B"), ("C", "F"), ("H", "I"), ("J", "K")):
            if f"{a}{r}:{b}{r}" not in [str(m) for m in ws.merged_cells.ranges]:
                ws.merge_cells(f"{a}{r}:{b}{r}")

    # styles for the rows that were swallowed by the giant merge
    for r in range(26, 54):
        for col in ("H", "I", "J", "K"):
            ws[f"{col}{r}"]._style = copy(ws[f"{col}25"]._style)

    # --- 2. per-line formulas
    ws["J7"] = "LARGO A PEDIR"
    for r in range(DETAIL_FIRST, DETAIL_LAST + 1):
        # L lee el precio directo de la columna C: si alguien inserta un
        # producto nuevo sin copiar la formula de la columna E, igual avisa
        # SIN PRECIO en vez de cotizar a 0.
        ws[f"L{r}"] = (
            f'=IF($C{r}="","",IFERROR(IF(VLOOKUP($C{r},PRECIOS!$A$2:$E${LAST_CAT},3,FALSE)="",'
            f'"SIN PRECIO",VLOOKUP($C{r},PRECIOS!$A$2:$E${LAST_CAT},3,FALSE)),"NO EN LISTA"))'
        )
        # J = largo comercial: se redondea hacia arriba al paso de corte.
        # Un "4.2" escrito con punto se guarda como fecha (46057) y rompería el
        # total, así que fuera del rango físico se avisa en vez de calcular.
        ws[f"J{r}"] = (
            f'=IF($C{r}="","",IFERROR(IF(VLOOKUP($C{r},PRECIOS!$A$2:$E${LAST_CAT},4,FALSE)<>"SI","",'
            f'IF($H{r}="","",IF(OR(NOT(ISNUMBER($H{r})),$H{r}<=0,$H{r}>50),"REVISAR LARGO",'
            f'ROUND(ROUNDUP($H{r}/PRECIOS!$H$1,0)*PRECIOS!$H$1,2)))),""))'
        )
        ws[f"M{r}"] = (
            f'=IF(OR($C{r}="",NOT(ISNUMBER($L{r}))),"",'
            f'IF(VLOOKUP($C{r},PRECIOS!$A$2:$E${LAST_CAT},4,FALSE)="SI",'
            f'IF(NOT(ISNUMBER($J{r})),"",$A{r}*$J{r}*$L{r}),$A{r}*$L{r}))'
        )
        ws[f"H{r}"].number_format = "0.00"
        ws[f"J{r}"].number_format = "0.00"
        ws[f"L{r}"].number_format = MONEY
        ws[f"M{r}"].number_format = MONEY

    # --- 3. product dropdown from the catalog
    keep = [d for d in ws.data_validations.dataValidation if "C8" not in str(d.sqref)]
    ws.data_validations.dataValidation = keep
    dv = DataValidation(
        type="list",
        formula1=f"PRECIOS!$A$2:$A${LAST_CAT}",
        allow_blank=True,
    )
    dv.error = "Elegí un producto de la lista."
    dv.errorTitle = "Producto no encontrado"
    dv.prompt = "Elegí el producto"
    dv.promptTitle = "PRODUCTO"
    dv.add(f"C{DETAIL_FIRST}:C{DETAIL_LAST}")
    ws.add_data_validation(dv)

    # el largo se escribe con coma decimal; con punto Excel lo toma como fecha
    dvl = DataValidation(
        type="decimal",
        operator="between",
        formula1="0.5",
        formula2="50",
        allow_blank=True,
    )
    dvl.errorTitle = "Largo invalido"
    dvl.error = (
        "Escribi el largo en metros con COMA decimal, por ejemplo 4,2 "
        "(con punto Excel lo lee como fecha). Entre 0,5 y 50 m."
    )
    dvl.promptTitle = "LARGO REAL"
    dvl.prompt = "Largo que pide el cliente, en metros: 4,2"
    dvl.add(f"H{DETAIL_FIRST}:H{DETAIL_LAST}")
    ws.add_data_validation(dvl)

    # --- 4. catalog sheet (hidden)
    sh = wb.create_sheet("PRECIOS")
    head = ["PRODUCTO", "UNIDAD", "PRECIO", "MULTIPLICA LARGO", "PRECIO OK"]
    thin = Side(style="thin", color="FFBFBFBF")
    for i, h in enumerate(head, 1):
        c = sh.cell(1, i, h)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor="FFD9E2F3")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(horizontal="center")
    for i, item in enumerate(catalog, 2):
        sh.cell(i, 1, item["name"])
        sh.cell(i, 2, item["unit"])
        p = sh.cell(i, 3, item["price"])
        p.number_format = MONEY
        sh.cell(i, 4, item["mult"])
        sh.cell(i, 5, f'=IF($C{i}="","SIN PRECIO",$C{i})').number_format = MONEY
        for col in range(1, 6):
            sh.cell(i, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sh.cell(i, 4).alignment = Alignment(horizontal="center")
    sh.column_dimensions["A"].width = 46
    sh.column_dimensions["B"].width = 16
    sh.column_dimensions["C"].width = 14
    sh.column_dimensions["D"].width = 18
    sh.column_dimensions["E"].width = 14
    sh.freeze_panes = "A2"
    sh.sheet_state = "hidden"

    # corte: los largos se piden en múltiplos de este paso
    sh["G1"] = "PASO DE CORTE (m)"
    sh["G1"].font = Font(bold=True, size=10)
    sh["G1"].fill = PatternFill("solid", fgColor="FFFFE699")
    sh["G1"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sh["H1"] = 0.5
    sh["H1"].number_format = "0.00"
    sh["H1"].fill = PatternFill("solid", fgColor="FFFFF2CC")
    sh["H1"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sh.column_dimensions["G"].width = 20

    # --- 5. footer total
    ws["C86"] = "TOTAL"
    ws["C86"].font = Font(bold=True, size=12, name="Calibri")
    ws["C86"].alignment = Alignment(horizontal="right")
    ws["L86"] = f"=SUM(M{DETAIL_FIRST}:M{DETAIL_LAST})"
    ws["L86"].number_format = MONEY
    ws["L86"].font = Font(bold=True, size=12, name="Calibri")

    wb.calculation.fullCalcOnLoad = True

    dst = os.path.join(OUT, "PRESUPUESTO.xlsx")
    wb.save(dst)
    inject_logo(PRESU, dst)
    return dst


def inject_logo(src, dst):
    """openpyxl drops drawings; put the original logo parts back."""
    with zipfile.ZipFile(src) as z:
        parts = {
            "xl/drawings/drawing1.xml": z.read("xl/drawings/drawing1.xml"),
            "xl/drawings/_rels/drawing1.xml.rels": z.read(
                "xl/drawings/_rels/drawing1.xml.rels"
            ),
            "xl/media/image1.png": z.read("xl/media/image1.png"),
            "xl/media/image2.jpg": z.read("xl/media/image2.jpg"),
        }
    sheet_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>'
    )
    tmp = dst + ".tmp"
    with zipfile.ZipFile(dst) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                assert "<drawing" not in text, "sheet already has a drawing"
                if "xmlns:r=" not in text.split(">", 1)[0]:
                    text = text.replace(
                        "<worksheet ",
                        '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
                        '2006/relationships" ',
                        1,
                    )
                    assert "xmlns:r=" in text.split(">", 1)[0], "no se pudo declarar xmlns:r"
                # CT_Worksheet: <drawing> comes after pageSetup/headerFooter and
                # before legacyDrawing/tableParts/extLst.
                for anchor in ("<legacyDrawing", "<tableParts", "<extLst"):
                    if anchor in text:
                        text = text.replace(anchor, '<drawing r:id="rId1"/>' + anchor, 1)
                        break
                else:
                    text = text.replace("</worksheet>", '<drawing r:id="rId1"/></worksheet>')
                ET.fromstring(text.encode("utf-8"))
                data = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                add = ""
                if 'Extension="png"' not in text:
                    add += '<Default Extension="png" ContentType="image/png"/>'
                if 'Extension="jpg"' not in text:
                    add += '<Default Extension="jpg" ContentType="image/jpeg"/>'
                if "/xl/drawings/drawing1.xml" not in text:
                    add += (
                        '<Override PartName="/xl/drawings/drawing1.xml" ContentType='
                        '"application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                    )
                data = text.replace("</Types>", add + "</Types>").encode("utf-8")
            zout.writestr(item, data)
        for name, data in parts.items():
            zout.writestr(name, data)
        zout.writestr("xl/worksheets/_rels/sheet1.xml.rels", sheet_rels)
    os.replace(tmp, dst)


# ------------------------------------------------------------------- INGRESOS
def build_ingresos():
    wb = openpyxl.load_workbook(INGRE)
    ws = wb["VENTAS DIARIAS"]
    # one sale has the date typed as text ("1/06//2026"); it never matches a
    # real date, so it silently disappears from any balance.
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.strip():
            ws.cell(r, 2).value = datetime(2026, 6, 1)
            ws.cell(r, 2).number_format = "dd/mm/yyyy"
            print(f"  fecha de texto corregida en fila {r}: {v!r} -> 01/06/2026")

    ws = wb["BALANCE"]
    rng = "'VENTAS DIARIAS'!$B$2:$B$5000"
    sums = "'VENTAS DIARIAS'!$E$2:$E$5000"
    n = 0
    for r in range(4, 38):
        a = ws[f"A{r}"].value
        if getattr(a, "year", None):
            ws[f"C{r}"] = f'=IF($A{r}="","",SUMIF({rng},$A{r},{sums}))'
            ws[f"C{r}"].number_format = MONEY
            n += 1

    # pre-existing template bug: the last week's totals only covered 3 of the 6
    # days (C32+C33+C34) and D..G were missing entirely.
    for col in "CDEFG":
        ref = ws[f"{col}31"]
        ws[f"{col}38"] = f"=SUM({col}32:{col}37)"
        ws[f"{col}38"]._style = copy(ref._style)
        ws[f"{col}38"].number_format = MONEY
    wb.calculation.fullCalcOnLoad = True
    dst = os.path.join(OUT, "Ingresos Y Gastos Semanales.xlsx")
    wb.save(dst)
    return dst, n


# -------------------------------------------------------------- LISTA PRECIOS
def build_lista():
    wb = openpyxl.load_workbook(LISTA)
    ws = wb["Hoja1"]

    # canonical names for the existing rows
    for r in range(2, 41):
        v = ws.cell(r, 2).value
        if v:
            ws.cell(r, 2).value = catmod.norm(v)

    if "B41:D41" in [str(m) for m in ws.merged_cells.ranges]:
        ws.unmerge_cells("B41:D41")
    footer_text = ws["B41"].value
    footer_style = {c: copy(ws[f"{c}41"]._style) for c in ("A", "B", "C", "D")}

    first_new = 41
    for i, item in enumerate(c for c in catalog if c["origin"] == "nuevo"):
        r = first_new + i
        for col in ("A", "B", "C", "D"):
            ws[f"{col}{r}"]._style = copy(ws[f"{col}40"]._style)
        ws.row_dimensions[r].height = ws.row_dimensions[40].height
        ws[f"B{r}"] = item["name"]
        if item["price"] is not None:
            ws[f"C{r}"] = item["price"]
        else:
            ws[f"A{r}"] = "SIN PRECIO"
        unit = {"METRO LINEAL": "Metro Lineal", "M2": "Metro 2"}.get(item["unit"], item["unit"].title())
        ws[f"D{r}"] = unit

    last = first_new + len([c for c in catalog if c["origin"] == "nuevo"])
    for col in ("A", "B", "C", "D"):
        ws[f"{col}{last}"]._style = footer_style[col]
    ws.row_dimensions[last].height = ws.row_dimensions[41].height
    ws[f"B{last}"] = footer_text
    ws.merge_cells(f"B{last}:D{last}")

    dst = os.path.join(OUT, "LISTA DE PRECIOS.xlsx")
    wb.save(dst)
    return dst, last


if __name__ == "__main__":
    p = build_presupuesto()
    print("presupuesto ->", p)
    i, n = build_ingresos()
    print("ingresos ->", i, "formulas:", n)
    l, last = build_lista()
    print("lista ->", l, "footer row:", last)
