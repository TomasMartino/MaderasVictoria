"""Functional proof: evaluate formulas and prove a new sale flows into the balance."""
import os
import shutil
import datetime

import formulas
import openpyxl

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "calc")
os.makedirs(TMP, exist_ok=True)

results = []


def check(cond, msg):
    results.append((bool(cond), msg))


# ============================== 1. PRESUPUESTO ==============================
presu = os.path.join(TMP, "PRESUPUESTO.xlsx")
shutil.copy2(os.path.join(OUT, "PRESUPUESTO.xlsx"), presu)
wb = openpyxl.load_workbook(presu)
ws = wb["CLIENTES"]
ws["A8"], ws["C8"], ws["H8"] = 10, "ROLLIZO 13-15CM", 3.66
ws["A9"], ws["C9"] = 4, "PUNTALES 13-15"
ws["A10"], ws["C10"] = 2, "ALAMO 1X4"
ws["C11"] = "INVENTADO XYZ"
ws["A12"], ws["C12"], ws["H12"] = 1, "ROLLIZO 13-15CM", 4.2
ws["A13"], ws["C13"], ws["H13"] = 1, "ROLLIZO 13-15CM", 4.7
ws["A14"], ws["C14"], ws["H14"] = 1, "ROLLIZO 13-15CM", 4.5
ws["A15"], ws["C15"], ws["H15"] = 2, "VIGA 23-25CM", 46057   # "4.2" con punto = fecha
# producto por metro lineal que antes NO multiplicaba el largo: se le carga
# precio a mano para probar que ahora si lo hace
ws["A16"], ws["C16"], ws["H16"] = 2, "VIGA 22X22CM", 3.1
wb["PRECIOS"]["C41"] = 9000
wb["PRECIOS"]["E41"] = None   # fila insertada a mano: sin la formula de PRECIO OK
wb.save(presu)


def solve(path, sheet, cells):
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    name = os.path.basename(path).upper()
    out = {}
    for cell in cells:
        key = f"'[{name}]{sheet.upper()}'!{cell}"
        for k, v in sol.items():
            if k.upper() == key:
                try:
                    out[cell] = v.value[0, 0]
                except Exception:
                    out[cell] = v
                break
    return out


r = solve(presu, "CLIENTES",
          ["L8", "J8", "M8", "L9", "J9", "M9", "L10", "M10", "L11", "M11",
           "J12", "M12", "J13", "M13", "J14", "M14", "J15", "M15", "L16", "J16", "M16", "L86"])
check(abs(r["L8"] - 5000) < 0.01, f"precio de lista rollizo: {r['L8']}")
check(abs(r["J8"] - 4.0) < 0.01, f"3,66 m se pide en 4,00: {r['J8']}")
check(abs(r["M8"] - 200000) < 0.01, f"10 un x 4,00 m x 5.000 = 200.000: {r['M8']}")
check(abs(r["L9"] - 5500) < 0.01, f"precio unitario puntales: {r['L9']}")
check(r["J9"] == "", f"puntales no lleva largo: {r['J9']!r}")
check(abs(r["M9"] - 22000) < 0.01, f"4 puntales = 22.000: {r['M9']}")
check(r["L10"] == "SIN PRECIO", f"producto sin precio avisa: {r['L10']}")
check(r["M10"] == "", f"sin precio no calcula total: {r['M10']!r}")
check(r["L11"] == "NO EN LISTA", f"texto libre avisa: {r['L11']}")
check(abs(r["J12"] - 4.5) < 0.01, f"largo real 4,2 -> se pide 4,50: {r['J12']}")
check(abs(r["M12"] - 22500) < 0.01, f"4,50 x 5.000 = 22.500: {r['M12']}")
check(abs(r["J13"] - 5.0) < 0.01, f"largo real 4,7 -> se pide 5,00: {r['J13']}")
check(abs(r["M13"] - 25000) < 0.01, f"5,00 x 5.000 = 25.000: {r['M13']}")
check(abs(r["J14"] - 4.5) < 0.01, f"largo real 4,5 -> se pide 4,50: {r['J14']}")
check(abs(r["M14"] - 22500) < 0.01, f"4,50 x 5.000 = 22.500: {r['M14']}")
check(r["J15"] == "REVISAR LARGO", f"largo escrito con punto se frena: {r['J15']!r}")
check(r["M15"] == "", f"y no calcula total: {r['M15']!r}")
check(abs(r["J16"] - 3.5) < 0.01, f"viga 22x22 (metro lineal) pide 3,50: {r['J16']}")
check(abs(r["L16"] - 9000) < 0.01, f"precio sin la formula de PRECIO OK: {r['L16']}")
check(abs(r["M16"] - 63000) < 0.01, f"2 x 3,50 x 9.000 = 63.000: {r['M16']}")
check(abs(r["L86"] - 355000) < 0.01, f"TOTAL del presupuesto: {r['L86']}")

# ============================== 2. BALANCE ==================================
ingre = os.path.join(TMP, "INGRESOS.xlsx")
shutil.copy2(os.path.join(OUT, "Ingresos Y Gastos Semanales.xlsx"), ingre)
cells = ["C7", "C8", "C11", "C12", "C14", "C23", "C27", "C10", "C17", "C24", "C31", "C38", "C39",
         "F39", "J39", "H39", "I39"]
r = solve(ingre, "BALANCE", cells)
for cell, exp in [("C7", 0), ("C8", 0), ("C11", 246000), ("C12", 745000), ("C14", 1030000),
                  ("C23", 84000), ("C27", 425500), ("C10", 0), ("C17", 2021000),
                  ("C24", 84000), ("C31", 425500), ("C38", 0), ("C39", 2530500)]:
    check(abs(r[cell] - exp) < 0.01, f"balance {cell} = {exp}: {r[cell]}")
check(abs(r["F39"] - 2500000) < 0.01, f"sueldos del mes 2.500.000: {r['F39']}")
check(abs(r["I39"] - 2500000) < 0.01, f"egresos del mes 2.500.000: {r['I39']}")
check(abs(r["J39"] - 30500) < 0.01, f"resultado neto del mes 30.500: {r['J39']}")

# ---- prove a NEW sale row flows into the balance automatically
wb2 = openpyxl.load_workbook(ingre)
wsv = wb2["VENTAS DIARIAS"]
last = wsv.max_row
wsv.cell(last + 1, 2, datetime.datetime(2026, 10, 22))
wsv.cell(last + 1, 3, "EFECTIVO")
wsv.cell(last + 1, 4, "PRUEBA")
wsv.cell(last + 1, 5, 111000)
wsv.cell(last + 2, 2, datetime.datetime(2026, 10, 22))
wsv.cell(last + 2, 3, "TRANSFERENCIA")
wsv.cell(last + 2, 4, "PRUEBA 2")
wsv.cell(last + 2, 5, 39000)
wb2.save(ingre)
r2 = solve(ingre, "BALANCE", ["C28", "C31", "C39"])
check(abs(r2["C28"] - 150000) < 0.01, f"venta nueva del 22/10 aparece sola: {r2['C28']}")
check(abs(r2["C31"] - 575500) < 0.01, f"total semanal se actualiza: {r2['C31']}")
check(abs(r2["C39"] - 2680500) < 0.01, f"total mensual se actualiza: {r2['C39']}")

bad = [m for ok, m in results if not ok]
for ok, m in results:
    print(("OK  " if ok else "MAL ") + m)
print(f"\n{len(results)-len(bad)}/{len(results)} OK")
if bad:
    raise SystemExit(1)
