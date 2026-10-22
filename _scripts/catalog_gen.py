"""Build the canonical product catalog shared by PRESUPUESTO and LISTA DE PRECIOS."""
import re
import openpyxl

BASE = r"C:\Users\User\Desktop\Proyecos\TioVity\respaldo 2026-10-22"
LISTA = BASE + r"\LISTA DE PRECIOS.xlsx"
PRESU = BASE + r"\PRESUPUESTO.xlsx"

# Names that arrive glued in the current dropdown.
RENAMES = {
    "VIGA22X22CM": "VIGA 22X22CM",
    "VIGA18X18CM": "VIGA 18X18CM",
    "ROLLIZO10X10CM": "ROLLIZO 10X10CM",
}


def norm(s):
    s = str(s).replace('""', '"')          # Excel doubles a literal quote in list validations
    s = re.sub(r"\s+", " ", s).strip().upper()
    return RENAMES.get(s, s)


# --- current price list (source of names/units/prices) -------------------
wb = openpyxl.load_workbook(LISTA)
ws = wb["Hoja1"]
list_rows = []  # (row, name, price, unit)
for r in range(2, 41):
    name = ws.cell(r, 2).value
    if not name:
        continue
    list_rows.append((r, norm(name), ws.cell(r, 3).value, ws.cell(r, 4).value))

list_names = {n for _, n, _, _ in list_rows}

# --- dropdown items ------------------------------------------------------
wp = openpyxl.load_workbook(PRESU)
wsp = wp["CLIENTES"]
dv = wsp.data_validations.dataValidation[0]
raw = dv.formula1.replace('"&"', "").strip('"')
dropdown = [norm(x) for x in raw.split(",") if x.strip()]

# --- explicit mapping: dropdown name -> existing list name ---------------
MAPPING = {
    "ROLLIZO 13/15CM": "ROLLIZO 13-15CM",
    "POSTES 16/18 CM": "POSTE 16-18CM",
    "VIGA 20/22 CM": "VIGA 20-22 CM",
    "MACHIMBRE 3/4 X4": "MACHIMBRE 3/4 X4 1RA",
    "FENOLICO 18MM 1 CB": "FENOLICOS 18MM 1 CARA BUENA",
    "VIGA ESPECIAL 23/25 CM": "VIGA 23-25CM",
    "MEMBRANA 40KG": "MEMBRANA 40KG",
    "PUNTALES 10/12CM": "PUNTALES 10-12",
    "IMPREGNADOS 10/12": "IMPREGNADOS 10-12",
    "FENOLICOS COMERCIAL18MM": "FENOLICOS 18MM COMERCIAL",
    "RUBEROID 20": "RUBEROID 20M",
    "TABLAS 3/4 X 4": "TABLAS DE PINO CEP. 3/4X4 3,66",
    "ALFAJIAS 1/2X2": "ALFAJIAS 1/2X2 2,20",
    "ALFAJIAS 2X1": "ALFAJIAS 2X1 2.20",
    "POSTES IMPREG. 16/18CM": "POSTES IMPREGNADO 16-18",
    "PUNTALES 13/15CM": "PUNTALES 13-15",
    "RUBEROID 40": "RUBEROID 40M",
    "ALFAJIAS 2X2": "ALFAJIAS 2X2 2,20",
    'MACHIMBRE 1"': 'MACHIMBRE 1"X6',
    "DECK PINO IMPREGNADO": "DECK PINO IMPREGNADO",
    "IMPREGNADOS 7/10": "IMPREGNADOS 7-10",
    "FENOLICO INDUSTRIAL 15MM": "FENOLICOS 15MM INDUSTRIAL",
    "MACHIMBRE 3/4 X 5": "MACHIMBRE 3/4 X5 1RA",
    "PUNTALES 7/10CM": "PUNTALES 7-10",
    "IMPREGNADO 13/15CM": "IMPREGNADOS 13-15",
    "VIGA IMPREGNADA 20/22CM": "VIGAS IMPREGNADAS 20-22",
    "FENOLICO COMERCIAL 16MM": "FENOLICOS 16MM COMERCIAL",
    "VIGA ESPECIAL 30CM": "VIGA 25-30CM",
    "MACHIMBRE 1/2 X 4": "MACHIMBRE 1/2",
}

missing_targets = [k for k, v in MAPPING.items() if v not in list_names]
assert not missing_targets, f"mapping targets not in price list: {missing_targets}"
unmapped_dropdown = [d for d in dropdown if d not in MAPPING]

# new products = unmapped dropdown names, deduped
new_names = []
for d in unmapped_dropdown:
    if d not in new_names:
        new_names.append(d)

# --- units ---------------------------------------------------------------
LARGO_KEYS = ("ROLLIZO", "POSTE", "VIGA")


def unit_for(name):
    n = name.upper()
    if "ALFAJIA" in n:
        return "UNIDAD"
    if "MACHIMBRE" in n or "DECK" in n or "FENOLICO" in n:
        return "METRO 2"
    if any(k in n for k in LARGO_KEYS):
        return "METRO LINEAL"
    return "UNIDAD"


def largo_flag(unit):
    """A price per linear meter is always multiplied by the length.

    Derived from the unit, never from the name: a product sold by "METRO
    LINEAL" that does not multiply the length undercharges by the whole
    length (a 4,5 m viga would be charged as 1 m).
    """
    return "SI" if unit == "METRO LINEAL" else "NO"


def numeric(v):
    return v if isinstance(v, (int, float)) else None


# --- catalog -------------------------------------------------------------
catalog = []  # dict(name, unit, price, mult, origin)
for _, name, price, unit in list_rows:
    u = norm(unit or "") or unit_for(name)
    if u == "M2":
        u = "METRO 2"
    catalog.append(
        dict(name=name, unit=u, price=numeric(price), mult=largo_flag(u), origin="lista")
    )

for name in new_names:
    u = unit_for(name)
    catalog.append(
        dict(name=name, unit=u, price=None, mult=largo_flag(u), origin="nuevo")
    )

# dedupe by name
seen = {}
for c in catalog:
    seen.setdefault(c["name"], c)
catalog = list(seen.values())

if __name__ == "__main__":
    print(f"lista={len(list_rows)} dropdown={len(dropdown)} mapeados={len(MAPPING)} nuevos={len(new_names)}")
    print(f"CATALOGO FINAL: {len(catalog)} productos\n")
    print(f"{'#':>3} {'PRODUCTO':<42} {'UNIDAD':<13} {'PRECIO':>11} {'LARGO':<5} ORIGEN")
    for i, c in enumerate(catalog, 1):
        p = "" if c["price"] is None else f"{c['price']:,.0f}"
        print(f"{i:>3} {c['name']:<42} {c['unit']:<13} {p:>11} {c['mult']:<5} {c['origin']}")
    con = sum(1 for c in catalog if c["price"] is not None)
    print(f"\ncon precio: {con} | sin precio: {len(catalog)-con}")

    # invariants
    names = [c["name"] for c in catalog]
    assert len(names) == len(set(names)), "duplicate catalog names: " + str(
        [n for n in names if names.count(n) > 1]
    )
    reachable = set(names) | set(MAPPING.values())
    lost = [d for d in dropdown if d not in reachable and MAPPING.get(d) not in reachable]
    print("dropdown no representado:", lost)
    assert not lost, lost

    bad = [c for c in catalog if (c["unit"] == "METRO LINEAL") != (c["mult"] == "SI")]
    assert not bad, "unidad y MULTIPLICA LARGO no coinciden: " + str(bad)
